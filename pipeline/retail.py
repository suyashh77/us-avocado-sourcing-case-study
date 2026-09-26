"""Extract clearly labeled retail price observations from USDA source files.

Requires openpyxl and the pdftotext executable (Poppler). The AMS URL is a
rolling latest-report PDF; delete the ignored raw PDF to refresh deliberately.
"""

import json
import re
import subprocess
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
STAGING = ROOT / "data" / "staging"


def scanner_benchmark():
    sheet = load_workbook(RAW / "avocado_retail_2023.xlsx", read_only=True, data_only=True).active
    row = next(row for row in sheet.iter_rows(values_only=True) if str(row[0]).startswith("Fresh1"))
    assert row[2].strip() == "per pound"
    price = float(row[1])
    assert 1 < price < 5
    return {"year": 2023, "usdPerPound": round(price, 4),
            "measure": "national scanner-data estimate of average retail price"}


def weekly_advertised():
    result = subprocess.run(["pdftotext", "-raw", str(RAW / "ams_retail_latest.pdf"), "-"],
                            check=True, capture_output=True, text=True, errors="replace")
    content = result.stdout
    date_match = re.search(r"Friday, ([A-Za-z]+ \d{1,2}, \d{4}) - Page 4", content)
    assert date_match, "AMS report date not found"
    report_date = datetime.strptime(date_match.group(1), "%B %d, %Y").date().isoformat()
    national = content.split("NATIONAL CONVENTIONAL SUMMARY", 1)[1].split("NORTHEAST U.S.", 1)[0]
    matches = re.findall(r"^Avocados Hass each\s+([\d,]+)\s+([\d.]+)\s+([\d,]+)\s+([\d.]+)\s+([\d,]+)\s+([\d.]+)\s*$", national, re.M)
    assert len(matches) == 1, "Expected one national conventional Hass-each price row"
    ads, price, prior_ads, prior, year_ago_ads, year_ago = matches[0]
    assert int(ads.replace(",", "")) > 1000 and 0 < float(price) < 5
    return {"reportDate": report_date, "product": "Hass, conventional, each",
            "advertisedUsdEach": float(price), "adCount": int(ads.replace(",", "")),
            "priorWeekUsdEach": float(prior), "priorWeekAdCount": int(prior_ads.replace(",", "")),
            "yearAgoUsdEach": float(year_ago), "yearAgoAdCount": int(year_ago_ads.replace(",", "")),
            "measure": "national weighted average weekly advertised price"}


if __name__ == "__main__":
    retail = {"scannerBenchmark": scanner_benchmark(), "weeklyAdvertised": weekly_advertised()}
    STAGING.mkdir(exist_ok=True)
    (STAGING / "retail_prices.json").write_text(json.dumps(retail, indent=2), encoding="utf-8")
    print(json.dumps(retail, indent=2))
